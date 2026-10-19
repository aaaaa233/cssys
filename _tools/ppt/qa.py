# -*- coding: utf-8 -*-
"""PPT 质检：文字转储 + 用本机 PowerPoint 导出 PNG（本机没装 LibreOffice）。

    python _tools/ppt/qa.py labs/lab7/幻灯片.pptx           # 只转储文字
    python _tools/ppt/qa.py labs/lab7/幻灯片.pptx --png     # 顺便导出图片
"""
import os
import sys
import shutil

from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400.0


def dump(path):
    prs = Presentation(path)
    W, H = prs.slide_width, prs.slide_height
    print("deck: %s   %d slides   %.3f x %.3f in"
          % (os.path.basename(path), len(prs.slides._sldIdLst),
             W / EMU_IN, H / EMU_IN))
    margin = 0.5 * EMU_IN
    # 页脚带：foot() 把页脚放在 H-0.75in，正文压到这里就会和页脚叠字。
    foot_top = H - 0.75 * EMU_IN
    for i, s in enumerate(prs.slides, 1):
        print("\n--- slide %d ---" % i)
        # 这一页到底有没有页脚？封面/节间/暗场页没有，就不该报页脚冲突。
        has_foot = False
        for sh in s.shapes:
            try:
                if sh.top is not None and abs(sh.top - foot_top) < 0.03 * EMU_IN:
                    has_foot = True
                    break
            except Exception:
                pass
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text.replace("\n", " / ")
                print("   %s" % t)
            if sh.has_table:
                for r in sh.table.rows:
                    print("   | " + " | ".join(
                        c.text_frame.text.strip() for c in r.cells))
            # 出血检查
            try:
                l, t_, w, h = sh.left, sh.top, sh.width, sh.height
            except Exception:
                continue
            if l is None:
                continue
            if l == 0 and t_ == 0 and w == W and h == H:
                continue                      # 整页背景，跳过
            kind = "TEXT" if sh.has_text_frame else "SHAPE"
            if l < margin - 1 or t_ < margin - 1 or l + w > W - margin + 1 \
                    or t_ + h > H - margin + 1:
                print("   !! %s 出界 [%.2f,%.2f  %.2fx%.2f]"
                      % (kind, l / EMU_IN, t_ / EMU_IN, w / EMU_IN, h / EMU_IN))
            # 页脚带占用：本身就在页脚位置的（页脚文字）不算
            elif has_foot and t_ + h > foot_top - 0.10 * EMU_IN \
                    and t_ < foot_top - 0.03 * EMU_IN:
                print("   !? %s 逼近页脚 底边 %.2f in（页脚在 %.2f）"
                      % (kind, (t_ + h) / EMU_IN, foot_top / EMU_IN))


def png(path, outdir):
    import win32com.client
    path = os.path.abspath(path)
    outdir = os.path.abspath(outdir)
    if os.path.isdir(outdir):
        shutil.rmtree(outdir)
    os.makedirs(outdir)
    app = win32com.client.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(path, ReadOnly=1, WithWindow=0)
    try:
        pres.Export(outdir, "PNG", 1600, 900)
    finally:
        pres.Close()
        app.Quit()
    files = sorted(os.listdir(outdir))
    print("\n%d png -> %s" % (len(files), outdir))
    for f in files:
        print("   %s" % os.path.join(outdir, f))


if __name__ == "__main__":
    p = sys.argv[1]
    dump(p)
    if "--png" in sys.argv:
        png(p, os.path.join(os.path.dirname(p), "_qa"))
